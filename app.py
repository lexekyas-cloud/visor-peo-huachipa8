# app.py
# -*- coding: utf-8 -*-

# ============== Importaciones ==============
import os, re, io
from datetime import datetime

import numpy as np
import pandas as pd
import streamlit as st
import plotly.express as px

from dotenv import load_dotenv
from st_aggrid import AgGrid, GridOptionsBuilder, GridUpdateMode, JsCode

from openpyxl import load_workbook
from openpyxl.styles import PatternFill


# ============== Configuración base ==============
st.set_page_config(page_title="Visor PMO - PEO Huachipa 8", layout="wide")
load_dotenv()

# ============== Banner superior (franja verde + logo TRUPAL) ==============
import base64, pathlib

APP_DIR = pathlib.Path(__file__).resolve().parent
logo_path = APP_DIR / "Logo_Blanco_Trupal.png"  # asegúrate de que el nombre coincida exactamente

def file_to_b64(p: pathlib.Path) -> str | None:
    try:
        with open(p, "rb") as f:
            return base64.b64encode(f.read()).decode()
    except Exception:
        return None

logo_b64 = file_to_b64(logo_path)

st.markdown(
    f"""
    <style>
      .top-banner {{
        background-color: #b5e61d; /* verde lima brillante */
        height: 70px;
        display: flex;
        align-items: center;
        justify-content: flex-start; /* usa center si lo quieres centrado */
        padding-left: 24px;
        border-bottom: 2px solid #6fbf6f;
      }}
      .top-banner img {{
        height: 55px;
      }}
      .banner-text {{
        font-family: 'Segoe UI', sans-serif;
        font-size: 26px;
        color: #1d4d1d;
        font-weight: bold;
        margin-left: 15px;
      }}
    </style>
    <div class="top-banner">
        {"<img src='data:image/png;base64," + logo_b64 + "' alt='TRUPAL'>" if logo_b64 else ""}
        <div class="banner-text">TRUPAL – Visor PMO</div>
    </div>
    """,
    unsafe_allow_html=True
)


# ============== fin bloque verde ==============

def _norm_path(p: str) -> str:
    """Normaliza una ruta proveniente del .env (quita comillas y corrige barras)."""
    if not p:
        return ""
    p = p.strip().strip('"').strip("'")
    return os.path.normpath(p)

EXCEL_PATH_RAW = os.getenv("EXCEL_PATH", "")
EXCEL_PATH     = _norm_path(EXCEL_PATH_RAW)
EXCEL_SHEET    = os.getenv("EXCEL_SHEET", "").strip() or 0  # nombre o índice

if not EXCEL_PATH or not os.path.exists(EXCEL_PATH):
    st.error(
        "⚠️ No encuentro el archivo Excel.\n\n"
        f"- EXCEL_PATH en .env: `{EXCEL_PATH_RAW}`\n"
        f"- Ruta normalizada: `{EXCEL_PATH}`\n\n"
        "Verifica que exista el archivo en esa ruta."
    )
    st.stop()


# ============== Helpers de conversión/estandarización ==============
def _coerce_date(x):
    if pd.isna(x) or x == "":
        return pd.NaT
    try:
        return pd.to_datetime(x, dayfirst=True, errors="coerce")
    except Exception:
        return pd.NaT

def _pct(x):
    if pd.isna(x) or x == "":
        return 0.0
    s = str(x).replace("%", "").replace(",", ".").strip()
    try:
        v = float(s)
        if 0 <= v <= 1:
            v *= 100
        return max(0.0, min(100.0, v))
    except Exception:
        return 0.0

def _derive_level_from_edt(edt):
    if pd.isna(edt):
        return np.nan
    s = str(edt).strip()
    return s.count(".") + 1 if re.match(r"^\d+(\.\d+)*$", s) else np.nan

def _safe_text(s):
    s = s.astype(str)
    s = s.str.replace(r"\.0$", "", regex=True)
    return s.replace("nan", "", regex=False)

def _standardize_cols(df):
    """Normaliza encabezados y tipos para que el resto del código sea robusto."""
    aliases = {
        "EDT":            ["edt", "wbs", "código edt", "codigo edt"],
        "Actividad":      ["actividad", "tarea", "nombre de tarea"],
        "Responsable":    ["responsable", "owner", "nombres de los recursos"],
        "FechaInicio":    ["inicio", "fechainicio", "comienzo", "fecha inicio"],
        "FechaFin":       ["fin", "fechafin", "fecha fin"],
        "PctAvance":      ["% avance", "% completado", "% completado real", "pctavance", "porcavance"],
        "PctTeorico":     ["% teorico", "% teórico", "%teorico", "%teórico", "pctteorico"],
        "Estado":         ["estado", "status"],
        "Evidencia":      ["evidencia", "link evidencia", "url"],
        "Comentario":     ["comentario", "observación", "observaciones"]
    }
    lower = {c.lower().strip(): c for c in df.columns}
    rename = {}
    for std, alts in aliases.items():
        for a in alts:
            if a in lower:
                rename[lower[a]] = std
                break
    df = df.rename(columns=rename)

    # Columnas mínimas
    for c in ["EDT", "Actividad", "Responsable", "FechaInicio", "FechaFin",
              "PctTeorico", "PctAvance", "Estado", "Evidencia", "Comentario"]:
        if c not in df:
            df[c] = np.nan

    # Tipos
    df["PctTeorico"]  = df["PctTeorico"].apply(_pct)
    df["PctAvance"]   = df["PctAvance"].apply(_pct)
    df["FechaInicio"] = df["FechaInicio"].apply(_coerce_date)
    df["FechaFin"]    = df["FechaFin"].apply(_coerce_date)

    # Nivel (si no viene)
    if "Nivel" not in df and "EDT" in df.columns:
        df["Nivel"] = df["EDT"].apply(_derive_level_from_edt)

    # Texto seguro
    for c in ["EDT", "Actividad", "Responsable", "Evidencia", "Comentario", "Estado"]:
        df[c] = _safe_text(df[c])

    return df

@st.cache_data(show_spinner=False)
def load_excel(path, sheet):
    df = pd.read_excel(path, sheet_name=sheet)
    return _standardize_cols(df)


def save_excel_with_highlight(original_df, updated_df, path):
    """
    Guarda updated_df en 'path' y pinta en MORADO las celdas que cambiaron
    respecto a original_df (comparando por fila/columna).
    """
    # Escribimos el DF actualizado
    updated_df.to_excel(path, index=False, engine="openpyxl")

    # Pintamos diferencias
    wb = load_workbook(path)
    ws = wb.active
    fill = PatternFill(start_color="EE82EE", end_color="EE82EE", fill_type="solid")  # morado

    common_cols = [c for c in updated_df.columns if c in original_df.columns]
    for r in range(len(updated_df)):
        for c, col in enumerate(common_cols, start=1):
            old = original_df.iloc[r][col] if r < len(original_df) else np.nan
            new = updated_df.iloc[r][col]
            if pd.isna(old) and pd.isna(new):
                continue
            if str(old) != str(new):
                ws.cell(row=r + 2, column=c).fill = fill

    wb.save(path)


# ============== JS para AgGrid ==============
GET_DATA_PATH_JS = JsCode("""
function (data) {
  if (!data || !data.EDT) return [];
  return String(data.EDT).split('.');
}
""")

INNER_RENDERER_JS = JsCode("""
function (p) { return p && p.data ? String(p.data.Actividad || '') : ''; }
""")

PCT_FORMATTER_JS = JsCode("""
function (p) {
  var v = (p && p.value != null) ? Math.round(p.value) : 0;
  return v + '%';
}
""")

# Estado con dropdown y regla de negocio que ajusta % avance
ESTADO_VALUESETTER_JS = JsCode("""
function(params){
  var v = params.newValue;
  params.data.Estado = v;
  if (v === 'Completada'){
    params.data.PctAvance = 100;
  } else if (v === 'Según lo programado'){
    var t = Number(params.data.PctTeorico || 0);
    params.data.PctAvance = t;
  }
  return true;
}
""")


# ============== Carga inicial de datos ==============
df_master = load_excel(EXCEL_PATH, EXCEL_SHEET)
df_show   = df_master.copy()

# Línea base (para Control de cambios)
if "baseline_df" not in st.session_state:
    st.session_state["baseline_df"] = df_master.copy()

# Fechas como texto para la grilla (no se guardan)
df_show["Inicio_txt"] = df_show["FechaInicio"].dt.strftime("%d/%m/%Y")
df_show["Fin_txt"]    = df_show["FechaFin"].dt.strftime("%d/%m/%Y")


# ============== Título y tabs ==============
st.title("📋 Visor PEO Huachipa 8")
tab1, tab2, tab3, tab4 = st.tabs(["📑 Actividades", "📊 Gantt", "📈 Dashboard", "📝 Control de cambios"])


# ============== TAB 1: ACTIVIDADES (edición directa) ==============
with tab1:
    st.caption(
        "Edita directamente **Estado**, **PctAvance**, **Evidencia** y **Comentario**. "
        "Al elegir *Completada* se pone 100%; con *Según lo programado* copia el % teórico."
    )

    # Columnas visibles
    cols = ["Responsable", "Inicio_txt", "Fin_txt", "PctTeorico", "PctAvance", "Estado", "Evidencia", "Comentario"]
    df_table = df_show[["EDT", "Actividad"] + cols].copy()

    g = GridOptionsBuilder.from_dataframe(df_table)
    g.configure_default_column(resizable=True, sortable=True, filter=True)

    # 🔹 Ocultamos EDT y Actividad (la jerarquía se ve en la columna agrupada)
    g.configure_column("EDT", hide=True)
    g.configure_column("Actividad", hide=True)

    # Editables
    g.configure_column("PctAvance", header_name="PctAvance", editable=True, valueFormatter=PCT_FORMATTER_JS)
    g.configure_column("PctTeorico", header_name="PctTeorico", editable=False, valueFormatter=PCT_FORMATTER_JS)
    g.configure_column("Evidencia", editable=True)
    g.configure_column("Comentario", editable=True)

    # Estado con dropdown + regla que ajusta % avance
    g.configure_column(
        "Estado",
        editable=True,
        cellEditor="agSelectCellEditor",
        cellEditorParams={"values": ["Retrasada", "Según lo programado", "Completada"]},
        valueSetter=ESTADO_VALUESETTER_JS
    )

    # Árbol por EDT (sin checkbox)
    g.configure_grid_options(
        treeData=True,
        getDataPath=GET_DATA_PATH_JS,
        groupDefaultExpanded=-1,
        animateRows=False,
    )
    g.configure_selection(selection_mode="single", use_checkbox=False)

    grid_options = g.build()
    grid_options["autoGroupColumnDef"] = {
        "headerName": "EDT / Actividad",
        "minWidth": 380,
        "cellRendererParams": {
            "suppressCount": True,
            "innerRenderer": INNER_RENDERER_JS,
            "checkbox": False
        }
    }

    grid = AgGrid(
        df_table,
        gridOptions=grid_options,
        update_mode=GridUpdateMode.VALUE_CHANGED,
        data_return_mode="AS_INPUT",
        allow_unsafe_jscode=True,
        theme="alpine",
        height=620,
        enable_enterprise_modules=True,
        key="grid_actividades"
    )

    # Sincronizamos lo editado hacia df_master por (EDT, Actividad)
    new_table = pd.DataFrame(grid["data"])  # columnas visibles
    if not new_table.empty:
        key = ["EDT", "Actividad"]
        df_sync = df_master.merge(
            new_table[key + ["PctAvance", "Estado", "Evidencia", "Comentario"]],
            on=key, how="left", suffixes=("", "_NEW")
        )
        for c in ["PctAvance", "Estado", "Evidencia", "Comentario"]:
            nc = f"{c}_NEW"
            if nc in df_sync.columns:
                df_sync[c] = np.where(df_sync[nc].notna(), df_sync[nc], df_sync[c])
                df_sync = df_sync.drop(columns=[nc])
        df_master = df_sync

    st.divider()
    colA, colB = st.columns(2)

    with colA:
        if st.button("💾 Guardar cambios en Excel", use_container_width=True, type="primary"):
            try:
                # Guardamos y pintamos diferencias
                original = load_excel(EXCEL_PATH, EXCEL_SHEET)
                save_excel_with_highlight(original, df_master, EXCEL_PATH)
                # Reseteamos línea base para Control de cambios
                st.session_state["baseline_df"] = df_master.copy()
                st.success("✅ Cambios guardados. Celdas cambiadas resaltadas en **morado** en el Excel.")
            except Exception as e:
                st.error(f"Error al guardar: {e}")

    with colB:
        buf = io.BytesIO()
        df_master.to_excel(buf, index=False, engine="openpyxl")
        buf.seek(0)
        st.download_button(
            "⬇️ Exportar Excel actualizado",
            data=buf,
            file_name=f"seguimiento_actualizado_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )


# ============== TAB 2: GANTT (simple, sin rojos) ==============
with tab2:
    st.subheader("📊 Gantt")
    st.caption("Gantt simple. Usa **FechaInicio** y **FechaFin**. (Paleta sin rojos).")

    gantt = df_master.copy()

    # Aseguramos tipo fecha
    for c in ["FechaInicio", "FechaFin"]:
        if c in gantt.columns:
            gantt[c] = pd.to_datetime(gantt[c], errors="coerce")

    # Selector de nivel (default 2 si existe)
    if "Nivel" in gantt.columns and gantt["Nivel"].notna().any():
        niveles = sorted(gantt["Nivel"].dropna().astype(int).unique().tolist())
        idx_default = niveles.index(2) if 2 in niveles else 0
        nivel_sel = st.selectbox("Nivel a mostrar", niveles, index=idx_default)
        gantt = gantt[gantt["Nivel"] == nivel_sel].copy()

    ok = gantt["FechaInicio"].notna() & gantt["FechaFin"].notna()
    gantt = gantt.loc[ok].copy()

    if gantt.empty:
        st.info("No hay tareas con **FechaInicio** y **FechaFin** para el nivel seleccionado.")
    else:
        edt_txt = gantt.get("EDT", "").fillna("").astype(str).str.replace(r"\.0$", "", regex=True)
        act_txt = gantt.get("Actividad", "").fillna("").astype(str)
        gantt["Tarea"] = edt_txt + " - " + act_txt

        palette_no_red = [
            "#1b9e77", "#66a61e", "#4daf4a", "#33a02c", "#7fc97f",
            "#1f78b4", "#80b1d3", "#a6cee3", "#2b8cbe", "#41ae76"
        ]
        color_col = "Responsable" if "Responsable" in gantt.columns else None

        fig = px.timeline(
            gantt,
            x_start="FechaInicio",
            x_end="FechaFin",
            y="Tarea",
            color=color_col,
            hover_data=[c for c in ["Estado", "PctAvance", "PctTeorico", "Evidencia", "Comentario", "Responsable"]
                        if c in gantt.columns],
            color_discrete_sequence=palette_no_red
        )
        fig.update_yaxes(autorange="reversed")
        fig.update_xaxes(dtick="M1", tickformat="%b %Y", side="top")
        fig.update_layout(height=600, margin=dict(l=10, r=10, t=30, b=10))
        st.plotly_chart(fig, use_container_width=True)


# ============== TAB 3: DASHBOARD Nivel 2 (con nombres, excluye Inicio 1.1) ==============
with tab3:
    st.subheader("📈 Dashboard – Atraso por Nivel 2")
    st.caption("Delay = **% Teórico − % Avance**. Valor positivo → *atraso*. Valor negativo → *adelantado*.")

    df = df_master.copy()

    # Pick robusto de columnas
    def pick_col(dfi, candidates):
        for c in candidates:
            if c in dfi.columns:
                return c
        return None

    teo_col = pick_col(df, ["PctTeorico", "PorcTeorico", "% Teórico", "% Teorico"])
    avc_col = pick_col(df, ["PctAvance", "PorcAvance", "% Avance", "% Completado"])

    if teo_col is None or avc_col is None:
        st.error("No encuentro columnas de %Teórico y/o %Avance (ej. 'PctTeorico' y 'PctAvance').")
        st.stop()

    for c in [teo_col, avc_col]:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0).clip(0, 100)

    # Helpers EDT
    def edt_parts(edt):
        if pd.isna(edt): return []
        s = str(edt).strip()
        return [p for p in s.split(".") if p != ""]

    def edt_prefix_2(edt):
        parts = edt_parts(edt)
        return ".".join(parts[:2]) if len(parts) >= 2 else None

    # Etiquetas Nivel 2
    label_df = pd.DataFrame(columns=["Nivel2", "ActividadN2"])
    if "Nivel" in df.columns and df["Nivel"].notna().any():
        tmp = df.copy()
        tmp["Nivel"] = pd.to_numeric(tmp["Nivel"], errors="coerce")
        if "EDT" in tmp.columns:
            lvl2_rows = tmp[(tmp["Nivel"] == 2) & tmp["EDT"].notna()].copy()
            if not lvl2_rows.empty:
                label_df = lvl2_rows[["EDT", "Actividad"]].rename(columns={"EDT": "Nivel2", "Actividad": "ActividadN2"})
    if label_df.empty and "EDT" in df.columns:
        lvl2_rows = df[df["EDT"].apply(lambda x: len(edt_parts(x)) == 2)].copy()
        if "Actividad" in lvl2_rows.columns and not lvl2_rows.empty:
            label_df = lvl2_rows[["EDT", "Actividad"]].rename(columns={"EDT": "Nivel2", "Actividad": "ActividadN2"})

    if "EDT" not in df.columns:
        st.error("No hay columna 'EDT' para agrupar por Nivel 2.")
        st.stop()

    df["Nivel2"] = df["EDT"].apply(edt_prefix_2)
    base = df[df["Nivel2"].notna()].copy()

    if base.empty:
        st.info("No hay datos a nivel 2 para armar el dashboard.")
        st.stop()

    # 🔹 Exclusiones personalizadas (hito Inicio 1.1)
    EXCLUIR_N2  = {"1.1"}
    EXCLUIR_ACT = {"Inicio"}

    base = base[~base["Nivel2"].isin(EXCLUIR_N2)]

    resumen = (
        base.groupby("Nivel2", as_index=False)
            .agg({teo_col: "mean", avc_col: "mean"})
            .rename(columns={teo_col: "%Teórico", avc_col: "%Avance"})
    )
    resumen["Delay"]     = (resumen["%Teórico"] - resumen["%Avance"]).round(1)
    resumen["%Teórico"]  = resumen["%Teórico"].round(1)
    resumen["%Avance"]   = resumen["%Avance"].round(1)

    if not label_df.empty:
        resumen = resumen.merge(label_df, on="Nivel2", how="left")

    if "ActividadN2" in resumen.columns:
        resumen = resumen[~resumen["ActividadN2"].fillna("").str.strip().isin(EXCLUIR_ACT)]

    resumen["Etiqueta"] = resumen["ActividadN2"].where(resumen["ActividadN2"].notna(), resumen["Nivel2"])
    resumen = resumen.sort_values("Delay", ascending=False).reset_index(drop=True)

    # KPIs
    colA, colB, colC = st.columns(3)
    atraso_prom  = resumen["Delay"].mean().round(1)
    n_atrasados  = int((resumen["Delay"] > 0.0).sum())
    n_total      = int(len(resumen))
    mejor        = resumen.loc[resumen["Delay"].idxmin(), "Etiqueta"] if not resumen.empty else "-"

    colA.metric("Atraso promedio (Nivel 2)", f"{atraso_prom} ptos")
    colB.metric("Unidades Nivel 2 con atraso", f"{n_atrasados} de {n_total}")
    colC.metric("Mejor Unidad (menor Delay)", f"{mejor}")

    st.divider()

    # Tabla resumen
    st.markdown("#### Tabla de atraso por Nivel 2")
    tabla = resumen[["Etiqueta", "%Teórico", "%Avance", "Delay"]].rename(columns={"Etiqueta": "Nivel 2"})
    st.dataframe(tabla, use_container_width=True, hide_index=True)

    # Teórico vs Avance
    st.markdown("#### % Teórico vs % Avance (por Nivel 2)")
    fig_tr = px.bar(
        resumen,
        x="Etiqueta", y=["%Teórico", "%Avance"],
        barmode="group",
        height=420,
        color_discrete_sequence=["#80b1d3", "#33a02c"]
    )
    fig_tr.update_layout(margin=dict(l=10, r=10, t=10, b=10), xaxis_title="Nivel 2 (Actividad)", yaxis_title="%")
    st.plotly_chart(fig_tr, use_container_width=True)

    # Delay
    st.markdown("#### Delay (Teórico − Avance)")
    fig_delay = px.bar(
        resumen.sort_values("Delay", ascending=False),
        x="Etiqueta", y="Delay",
        color="Delay",
        color_continuous_scale=["#1f78b4", "#c7c7c7", "#feb24c"],
        height=420
    )
    fig_delay.update_layout(margin=dict(l=10, r=10, t=10, b=10), xaxis_title="Nivel 2 (Actividad)", yaxis_title="ptos")
    st.plotly_chart(fig_delay, use_container_width=True)

    # Export CSV
    with st.expander("⬇️ Exportar datos del dashboard (CSV)"):
        csv = tabla.to_csv(index=False).encode("utf-8-sig")
        st.download_button("Descargar CSV", data=csv, file_name="dashboard_nivel2_con_nombres.csv", mime="text/csv")


# ============== TAB 4: CONTROL DE CAMBIOS (contra baseline) ==============
def _diff_changes(df_old, df_new, keys, observed_cols):
    """
    Compara df_old vs df_new y devuelve un dataframe con cambios en observed_cols.
    keys = columnas clave para identificar filas (EDT, Actividad, Responsable).
    """
    left  = df_old[keys + observed_cols].copy()
    right = df_new[keys + observed_cols].copy()
    merged = left.merge(right, on=keys, how="outer", suffixes=("_old", "_new"), indicator=True)

    rows = []
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for _, r in merged.iterrows():
        if r["_merge"] != "both":
            # si quieres registrar altas/bajas, manejar aquí
            continue
        for col in observed_cols:
            old = r.get(f"{col}_old", np.nan)
            new = r.get(f"{col}_new", np.nan)
            if (pd.isna(old) and pd.isna(new)) or str(old) == str(new):
                continue
            rows.append({
                "FechaCambio": now,
                "EDT": r.get("EDT", ""),
                "Actividad": r.get("Actividad", ""),
                "Responsable": r.get("Responsable", ""),
                "Columna": col,
                "Valor_Anterior": old,
                "Valor_Nuevo": new,
            })

    return pd.DataFrame(rows, columns=[
        "FechaCambio", "EDT", "Actividad", "Responsable", "Columna", "Valor_Anterior", "Valor_Nuevo"
    ])

with tab4:
    st.subheader("📝 Control de cambios (desde última carga/guardado)")
    st.caption("Compara con la **línea base** (al cargar o tras el último **Guardar en Excel**).")

    observed_cols = ["Estado", "PctAvance", "Evidencia", "Comentario"]
    keys = ["EDT", "Actividad", "Responsable"]

    baseline   = st.session_state.get("baseline_df", df_master.copy())
    changes_df = _diff_changes(baseline, df_master, keys, observed_cols)

    if changes_df.empty:
        st.success("No hay cambios detectados. ✅")
    else:
        st.dataframe(changes_df, use_container_width=True, hide_index=True)
        csv = changes_df.to_csv(index=False).encode("utf-8-sig")
        st.download_button(
            "⬇️ Descargar control de cambios (CSV)",
            data=csv,
            file_name=f"control_cambios_{datetime.now().strftime('%Y%m%d_%H%M')}.csv",
            mime="text/csv"
        )
